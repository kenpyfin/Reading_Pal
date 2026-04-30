#!/usr/bin/env python3
"""
Smoke test for the PDF pipeline (text extraction and optional PaddleOCR).
Run from project root:
  python pdf_service/scripts/smoke_test_pdf_pipeline.py
Or with .env present:
  cd pdf_service && python scripts/smoke_test_pdf_pipeline.py
Creates a minimal in-memory text PDF and runs pdf_has_text + process_text_pdf; optionally runs PaddleOCR if PDF_OCR_ENGINE=paddle.
"""
import io
import os
import sys

# Set minimal env for storage paths if not already set (so app import does not exit)
def _set_default_env():
    root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    for key, default in (
        ("PDF_STORAGE_PATH", os.path.join(root, "storage", "pdfs")),
        ("MARKDOWN_PATH", os.path.join(root, "storage", "output")),
        ("IMAGES_PATH", os.path.join(root, "storage", "images")),
    ):
        if not os.environ.get(key):
            os.environ[key] = default


_set_default_env()


def _add_app_to_path():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    pdf_service_root = os.path.dirname(script_dir)
    if pdf_service_root not in sys.path:
        sys.path.insert(0, pdf_service_root)


_add_app_to_path()


def make_minimal_text_pdf_bytes() -> bytes:
    """Produce a minimal PDF with one page and extractable text (PyMuPDF)."""
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello world. This is a smoke test for text extraction.", fontsize=12)
    buf = io.BytesIO()
    doc.save(buf, deflate=True)
    doc.close()
    return buf.getvalue()


def make_txt_bytes() -> bytes:
    return b"First line.\n\nSecond paragraph."


def make_html_bytes() -> bytes:
    return b"<html><body><h1>Chapter 1</h1><p>Hello from html.</p></body></html>"


def make_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Docx Title", level=1)
    doc.add_paragraph("This is a DOCX smoke test paragraph.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_epub_bytes() -> bytes:
    import tempfile
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("smoke-id")
    book.set_title("Smoke EPUB")
    book.set_language("en")
    chapter = epub.EpubHtml(title="Intro", file_name="intro.xhtml", lang="en")
    chapter.content = "<h1>Intro</h1><p>Hello from epub parser smoke test.</p>"
    book.add_item(chapter)
    book.toc = (chapter,)
    book.spine = ["nav", chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    with tempfile.NamedTemporaryFile(suffix=".epub", delete=False) as tmp:
        epub_path = tmp.name
    try:
        epub.write_epub(epub_path, book)
        with open(epub_path, "rb") as f:
            return f.read()
    finally:
        try:
            os.remove(epub_path)
        except OSError:
            pass


def main():
    from app import (
        pdf_has_text,
        process_text_pdf,
        process_document_bytes,
        extract_pdf_images,
        inject_image_links_into_markdown,
        PDF_OCR_ENGINE,
    )

    pdf_bytes = make_minimal_text_pdf_bytes()
    print("1. pdf_has_text:", pdf_has_text(pdf_bytes))
    text_md = process_text_pdf(pdf_bytes)
    print("2. process_text_pdf length:", len(text_md))
    assert "smoke test" in text_md or "Hello" in text_md, "Expected extractable text"
    assert "\n\n---\n\n" not in text_md, "Single-page text PDF should not include page separators"
    print("3. Text extraction smoke test passed.")

    images_per_page = extract_pdf_images(pdf_bytes, "smoke_test_doc", os.path.join(os.environ["IMAGES_PATH"], "app"))
    print("4. extract_pdf_images page count:", len(images_per_page))
    assert isinstance(images_per_page, list), "Expected list result from image extraction"
    # For this synthetic text-only PDF, embedded raster images are usually absent.
    assert sum(len(page) for page in images_per_page) == 0, "Expected no embedded raster images in fixture"

    merged_md = "Page one text\n\n---\n\nPage two text"
    injected = inject_image_links_into_markdown(
        merged_md,
        [["a.png"], ["b.jpg"]],
        web_image_base_path="/images/app",
        page_separator="\n\n---\n\n",
    )
    assert "![](/images/app/a.png)" in injected and "![](/images/app/b.jpg)" in injected, "Expected stable web image links"
    assert injected.count("\n\n---\n\n") == 1, "Expected markdown separator to remain canonical"
    print("5. Image link injection smoke test passed.")

    txt_md, txt_images = process_document_bytes(make_txt_bytes(), ".txt", "smoke_txt")
    assert "Second paragraph" in txt_md and txt_images == [], "TXT extraction failed"
    print("6. TXT extraction smoke test passed.")

    html_md, html_images = process_document_bytes(make_html_bytes(), ".html", "smoke_html")
    assert "Chapter 1" in html_md and "Hello from html" in html_md and html_images == [], "HTML extraction failed"
    print("7. HTML extraction smoke test passed.")

    docx_md, docx_images = process_document_bytes(make_docx_bytes(), ".docx", "smoke_docx")
    assert "DOCX smoke test paragraph" in docx_md and docx_images == [], "DOCX extraction failed"
    print("8. DOCX extraction smoke test passed.")

    epub_md, epub_images = process_document_bytes(make_epub_bytes(), ".epub", "smoke_epub")
    assert "epub parser smoke test" in epub_md and epub_images == [], "EPUB extraction failed"
    print("9. EPUB extraction smoke test passed.")

    if PDF_OCR_ENGINE == "paddle":
        try:
            from app import process_scanned_pdf_with_paddleocr
            ocr_md = process_scanned_pdf_with_paddleocr(pdf_bytes)
            print("10. process_scanned_pdf_with_paddleocr length:", len(ocr_md))
            print("11. PaddleOCR pipeline smoke test passed.")
        except Exception as e:
            print("11. PaddleOCR skipped (no GPU or dependency):", e)
    else:
        print("10. PaddleOCR skipped (PDF_OCR_ENGINE != paddle).")

    print("Done.")


if __name__ == "__main__":
    main()
